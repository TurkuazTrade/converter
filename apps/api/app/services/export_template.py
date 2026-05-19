from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from app.utils.normalization import normalize_key


SERVICE_FIELD_LABELS = {
    normalize_key("CLIENT CODE"): "client_code",
    normalize_key("DATE"): "document_date",
    normalize_key("WH NO"): "warehouse_no",
    normalize_key("FICHE NO"): "fiche_no",
}

HEADER_LABELS = {
    normalize_key("Unit"): "unit",
    normalize_key("Unit Price (Tenge)"): "unit_price",
}


@dataclass(frozen=True, slots=True)
class ExportTemplateLayout:
    sheet_name: str
    service_cells: dict[str, str]
    data_start_row: int
    item_code_col: int
    item_name_col: int | None
    unit_col: int | None
    quantity_col: int
    unit_price_col: int | None
    data_style_columns: tuple[int, ...]
    merged_ranges: tuple[str, ...]
    hidden_rows: tuple[int, ...]
    hidden_columns: tuple[str, ...]
    formulas: tuple[str, ...]
    active_columns: tuple[int, ...]

    @property
    def clearable_columns(self) -> tuple[int, ...]:
        columns = {
            self.item_code_col,
            self.quantity_col,
            *(col for col in (self.item_name_col, self.unit_col) if col is not None),
        }
        return tuple(sorted(columns))


def analyze_export_template(template: Path | str | Any) -> ExportTemplateLayout:
    close_workbook = False
    if isinstance(template, (Path, str)):
        workbook = openpyxl.load_workbook(template, data_only=False)
        close_workbook = True
    else:
        workbook = template

    try:
        worksheet = workbook.active
        service_cells, max_service_row = _service_cells(worksheet)
        headers = _header_cells(worksheet)
        data_start_row = _find_data_start_row(worksheet, max_service_row)
        sample_columns = _non_empty_columns(worksheet, data_start_row)
        if len(sample_columns) < 2:
            raise ValueError("Export template data row is not detectable.")

        unit_col = headers.get("unit")
        unit_price_col = headers.get("unit_price")
        item_code_col = _infer_item_code_col(sample_columns, unit_col)
        item_name_col = _infer_item_name_col(sample_columns, item_code_col, unit_col)
        quantity_col = _infer_quantity_col(worksheet, data_start_row, sample_columns, unit_col)
        data_style_columns = _data_style_columns(worksheet, data_start_row, sample_columns)

        return ExportTemplateLayout(
            sheet_name=worksheet.title,
            service_cells=service_cells,
            data_start_row=data_start_row,
            item_code_col=item_code_col,
            item_name_col=item_name_col,
            unit_col=unit_col,
            quantity_col=quantity_col,
            unit_price_col=unit_price_col,
            data_style_columns=tuple(data_style_columns),
            merged_ranges=tuple(str(item) for item in worksheet.merged_cells.ranges),
            hidden_rows=tuple(index for index, dim in worksheet.row_dimensions.items() if dim.hidden),
            hidden_columns=tuple(key for key, dim in worksheet.column_dimensions.items() if dim.hidden),
            formulas=tuple(_formula_cells(worksheet)),
            active_columns=tuple(sample_columns),
        )
    finally:
        if close_workbook:
            workbook.close()


def describe_export_template(template_path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(template_path, data_only=False)
    try:
        layout = analyze_export_template(workbook)
        worksheet = workbook[layout.sheet_name]
        return {
            "sheet_names": workbook.sheetnames,
            "active_sheet": layout.sheet_name,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "service_cells": layout.service_cells,
            "data_start_row": layout.data_start_row,
            "columns": {
                "item_code": get_column_letter(layout.item_code_col),
                "item_name": get_column_letter(layout.item_name_col) if layout.item_name_col else None,
                "unit": get_column_letter(layout.unit_col) if layout.unit_col else None,
                "quantity": get_column_letter(layout.quantity_col),
                "unit_price": get_column_letter(layout.unit_price_col) if layout.unit_price_col else None,
            },
            "data_style_columns": [get_column_letter(col) for col in layout.data_style_columns],
            "active_columns": [get_column_letter(col) for col in layout.active_columns],
            "merged_ranges": list(layout.merged_ranges),
            "hidden_rows": list(layout.hidden_rows),
            "hidden_columns": list(layout.hidden_columns),
            "formulas": list(layout.formulas),
            "column_widths": {
                get_column_letter(index): worksheet.column_dimensions[get_column_letter(index)].width
                for index in range(1, worksheet.max_column + 1)
            },
        }
    finally:
        workbook.close()


def _service_cells(worksheet: Any) -> tuple[dict[str, str], int]:
    service_cells: dict[str, str] = {}
    max_service_row = 0
    for row in worksheet.iter_rows():
        for cell in row:
            field = SERVICE_FIELD_LABELS.get(normalize_key(cell.value))
            if field is None:
                continue
            service_cells[field] = worksheet.cell(cell.row, cell.column + 1).coordinate
            max_service_row = max(max_service_row, cell.row)

    missing = sorted(set(SERVICE_FIELD_LABELS.values()) - set(service_cells))
    if missing:
        raise ValueError(f"Export template is missing service fields: {', '.join(missing)}.")
    return service_cells, max_service_row


def _header_cells(worksheet: Any) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            field = HEADER_LABELS.get(normalize_key(cell.value))
            if field is not None:
                result[field] = cell.column
    return result


def _find_data_start_row(worksheet: Any, max_service_row: int) -> int:
    for row_index in range(max_service_row + 1, worksheet.max_row + 1):
        columns = _non_empty_columns(worksheet, row_index)
        if len(columns) >= 2 and any(_has_data_format(worksheet.cell(row_index, col)) for col in columns):
            return row_index
    raise ValueError("Export template data start row was not found.")


def _non_empty_columns(worksheet: Any, row_index: int) -> list[int]:
    return [
        col
        for col in range(1, worksheet.max_column + 1)
        if worksheet.cell(row_index, col).value not in (None, "")
    ]


def _infer_item_code_col(sample_columns: list[int], unit_col: int | None) -> int:
    if unit_col is not None:
        before_unit = [col for col in sample_columns if col < unit_col]
        if before_unit:
            return before_unit[0]
    return sample_columns[0]


def _infer_item_name_col(sample_columns: list[int], item_code_col: int, unit_col: int | None) -> int | None:
    upper_bound = unit_col if unit_col is not None else max(sample_columns) + 1
    candidates = [col for col in sample_columns if item_code_col < col < upper_bound]
    return candidates[0] if candidates else None


def _infer_quantity_col(worksheet: Any, row_index: int, sample_columns: list[int], unit_col: int | None) -> int:
    numeric_columns = [
        col
        for col in sample_columns
        if col != unit_col and isinstance(worksheet.cell(row_index, col).value, (int, float))
    ]
    if numeric_columns:
        return numeric_columns[-1]
    if unit_col is not None:
        after_unit = [col for col in sample_columns if col > unit_col]
        if after_unit:
            return after_unit[-1]
    return sample_columns[-1]


def _data_style_columns(worksheet: Any, row_index: int, sample_columns: list[int]) -> list[int]:
    styled = [
        col
        for col in range(1, worksheet.max_column + 1)
        if _has_data_style(worksheet.cell(row_index, col))
    ]
    if styled:
        return styled
    return sample_columns


def _has_data_style(cell: Any) -> bool:
    return bool(cell.value not in (None, "") or _has_data_format(cell))


def _has_data_format(cell: Any) -> bool:
    border = cell.border
    has_border = any(
        side.style is not None for side in (border.left, border.right, border.top, border.bottom)
    )
    fill = cell.fill
    has_fill = fill.fill_type is not None
    return bool(has_border or has_fill)


def _formula_cells(worksheet: Any) -> list[str]:
    formulas: list[str] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.coordinate)
    return formulas
