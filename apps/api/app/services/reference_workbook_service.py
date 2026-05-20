from __future__ import annotations

from io import BytesIO
from typing import Iterable

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.models.client import Client
from app.models.product import Product
from app.services.converter_registry_service import ConverterRegistryService


class ReferenceWorkbookService:
    mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    product_headers = [
        "Номер товара",
        "Штрихкод",
        "Наименование",
        "Код сети",
        "Ценовой код",
        "Код обмена",
        "Артикул",
        "Остаток",
        "Торговая марка",
        "Бренд",
        "Тип",
        "Коэффициент",
    ]
    client_headers = [
        "Код клиента панорама",
        "Название клиента панорама",
        "Название клиента питон",
        "Адрес",
        "Сеть",
    ]

    def build_products_workbook(self, products: Iterable[Product] = ()) -> bytes:
        workbook = self._new_workbook("convert", self.product_headers)
        sheet = workbook.active
        for product in products:
            active_barcodes = [
                barcode.barcode
                for barcode in product.barcodes
                if barcode.deleted_at is None and barcode.is_active
            ]
            sheet.append(
                [
                    product.item_code,
                    ", ".join(active_barcodes),
                    product.name,
                    None,
                    product.price_code,
                    product.exchange_code,
                    product.article,
                    product.stock,
                    product.trade_mark,
                    product.brand,
                    product.product_type,
                    float(product.conversion_multiplier or 1),
                ]
            )
        return self._save(workbook)

    def build_clients_workbook(self, clients: Iterable[Client] = ()) -> bytes:
        workbook = self._new_workbook("client", self.client_headers)
        sheet = workbook.active
        for client in clients:
            sheet.append(
                [
                    client.client_code,
                    client.name,
                    client.name_2,
                    client.address,
                    client.network_name,
                ]
            )
        return self._save(workbook)

    def build_order_template(self, converter_type: str | None = None) -> bytes:
        config = ConverterRegistryService().load_config(converter_type or "asia_retail")
        columns = config.get("columns", {})
        headers = [
            self._first_alias(columns, "document_no", "Номер"),
            self._first_alias(columns, "document_date", "Дата"),
            self._first_alias(columns, "client_code", "Код клиента"),
            self._first_alias(columns, "client_name", "Клиент"),
            self._first_alias(columns, "client_branch", "Подразделение"),
            self._first_alias(columns, "item_code", "Код товара"),
            self._first_alias(columns, "item_name", "Товар"),
            self._first_alias(columns, "barcode", "Штрихкод"),
            self._first_alias(columns, "quantity", "Количество"),
        ]
        workbook = self._new_workbook(self._first_sheet_name(config), headers)
        sheet = workbook.active
        sheet.append(["ORDER-001", "2026-05-19", "100245", "Клиент", "Филиал", "ERP-100", "Товар", "4600000000000", 1])
        return self._save(workbook)

    @staticmethod
    def _new_workbook(sheet_name: str, headers: list[str]) -> openpyxl.Workbook:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = sheet_name[:31]
        sheet.append(headers)
        ReferenceWorkbookService._style_header(sheet)
        for index, header in enumerate(headers, start=1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(len(header) + 4, 16)
        sheet.freeze_panes = "A2"
        return workbook

    @staticmethod
    def _style_header(sheet: Worksheet) -> None:
        fill = PatternFill("solid", fgColor="E2E8F0")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = fill

    @staticmethod
    def _first_alias(columns: dict, field: str, fallback: str) -> str:
        value = columns.get(field)
        if isinstance(value, list) and value:
            return str(value[0])
        return fallback

    @staticmethod
    def _first_sheet_name(config: dict) -> str:
        aliases = config.get("sheet", {}).get("name_aliases") or ["Order"]
        return str(aliases[0] or "Order")

    @staticmethod
    def _save(workbook: openpyxl.Workbook) -> bytes:
        try:
            buffer = BytesIO()
            workbook.save(buffer)
            return buffer.getvalue()
        finally:
            workbook.close()
