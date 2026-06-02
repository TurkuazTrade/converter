from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal, ROUND_FLOOR
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.constants import DEFAULT_BRANCH_ID
from app.core.constants import CONVERTER_FILENAME_PREFIXES
from app.core.enums import OrderItemStatus, OrderStatus, ProcessingEventType
from app.models.order import Order, OrderItem, ProcessingEvent
from app.models.product import ProductTypeExportRule
from app.models.setting import AppSetting
from app.services.export_template import analyze_export_template
from app.utils.normalization import is_short_numeric_item_code, sanitize_filename_part


@dataclass(slots=True)
class ExportLine:
    item_code: str
    item_name: str | None
    quantity: float
    product_type: str | None = None
    row_number: int | None = None
    raw_item_code: str | None = None
    raw_barcode: str | None = None
    normalized_barcode: str | None = None
    warehouse_no: str | int | None = None
    unit: int | float = 1
    unit_price: float | None = None


@dataclass(slots=True)
class ExportProblemLine:
    row_number: int | None
    item_code: str | None
    item_name: str | None
    raw_item_code: str | None
    raw_barcode: str | None
    normalized_barcode: str | None
    raw_name: str | None
    product_type: str | None
    quantity: float | None
    reason: str


@dataclass(slots=True)
class ResolvedOrderExport:
    converter_type: str
    client_code: str
    document_date: date
    fiche_no: str
    lines: list[ExportLine]
    client_name: str | None = None
    order_id: int | None = None
    warehouse_no: str | int | None = None
    problems: list[ExportProblemLine] | None = None
    product_type_filter: str | None = None
    sequence_number: int = 0


@dataclass(slots=True)
class GeneratedExport:
    filename: str
    content: bytes
    mime_type: str
    size: int
    sha256: str


class ExportService:
    """Template-preserving export service."""

    FICHE_SEQUENCE_STEP = 5
    LEGACY_EXPORT_SEQUENCE_STRIDE = 50
    EXPORT_SEQUENCE_SETTING_KEY = "export_sequence_next"
    MAX_FICHE_SEQUENCE = 9_999_999_999

    def __init__(
        self,
        template_path: Path | None = None,
    ) -> None:
        self.template_path = template_path or settings.template_path

    def export_order(
        self,
        db: Session,
        order_id: int,
        user_id: int | None = None,
        *,
        product_type: str | None = None,
        branch_id: int | None = None,
    ) -> GeneratedExport:
        order = db.get(Order, order_id)
        if order is None:
            raise ValueError("Order not found.")
        if branch_id is not None and order.branch_id != branch_id:
            raise ValueError("Order not found.")
        sequence_number = self._next_export_sequence_number(db, branch_id=branch_id)
        payload = self._payload_from_order(
            order,
            sequence_number=sequence_number,
            excluded_product_types=self._excluded_product_types(db, branch_id=branch_id),
            product_type_filter=product_type,
        )
        export_sequence_count = self._export_sequence_count(payload)
        export_sequence_next = sequence_number + export_sequence_count * self.FICHE_SEQUENCE_STEP
        content = self.build_export_bytes(payload)
        filename = self.build_filename(payload)
        result = GeneratedExport(
            filename=filename,
            content=content,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size=len(content),
            sha256=sha256(content).hexdigest(),
        )
        order.status = OrderStatus.EXPORTED.value
        order.export_file_id = None
        order.error_message = None
        db.add(
            ProcessingEvent(
                order_id=order.id,
                event_type=ProcessingEventType.EXPORTED.value,
                message="Export workbook generated on demand.",
                payload={
                    "filename": result.filename,
                    "size": result.size,
                    "sha256": result.sha256,
                    "export_sequence_start": sequence_number,
                    "export_sequence_count": export_sequence_count,
                    "export_sequence_step": self.FICHE_SEQUENCE_STEP,
                    "export_sequence_next": export_sequence_next,
                },
                created_by_id=user_id,
            )
        )
        self._store_export_sequence_number(db, export_sequence_next, branch_id=branch_id)
        return result

    def build_export_bytes(self, order: ResolvedOrderExport) -> bytes:
        workbook = openpyxl.load_workbook(self.template_path)
        try:
            layout = analyze_export_template(workbook)
            template_sheet = workbook[layout.sheet_name]
            lines_by_type: dict[str, list[ExportLine]] = defaultdict(list)
            for line in order.lines:
                lines_by_type[self._display_product_type(line.product_type)].append(line)

            grouped_lines = sorted(lines_by_type.items(), key=lambda item: item[0].casefold())
            self._write_export_blocks_sheet(template_sheet, layout, order, grouped_lines)

            if order.problems:
                self._write_problem_sheet(workbook, order.problems)
            self._write_all_products_sheet(workbook, order, order.lines, order.problems or [])

            self._normalize_workbook_views(workbook)

            buffer = BytesIO()
            workbook.save(buffer)
            content = buffer.getvalue()
        finally:
            workbook.close()
        self.validate_export_bytes(content)
        return content

    def build_filename(self, order: ResolvedOrderExport) -> str:
        converter = self._filename_safe_text(
            self._filename_converter_prefix(order.converter_type),
            "Converter",
        )
        sequence = max(order.sequence_number, 0)
        order_id = f" id-{order.order_id}" if order.order_id is not None else ""
        suffix = ""
        if order.product_type_filter:
            suffix = f" {self._filename_safe_text(order.product_type_filter, 'type')}"
        client_name = (
            f" {self._filename_safe_text(order.client_name, 'client')}"
            if order.client_name and order.client_name.strip()
            else ""
        )
        export_date = order.document_date.isoformat()
        return f"{export_date}{order_id} {converter} zakaz {sequence:04d}{suffix}{client_name}.xlsx"

    @staticmethod
    def _filename_converter_prefix(converter_type: str) -> str:
        key = converter_type.strip().casefold()
        if key in CONVERTER_FILENAME_PREFIXES:
            return CONVERTER_FILENAME_PREFIXES[key]
        parts = [part for part in re.split(r"[^0-9A-Za-zА-Яа-я]+", converter_type.strip()) if part]
        if not parts:
            return "Converter"
        return "".join(part[:1].upper() + part[1:] for part in parts)

    @staticmethod
    def validate_export_bytes(content: bytes) -> None:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        workbook.close()

    def _payload_from_order(
        self,
        order: Order,
        *,
        sequence_number: int | None = None,
        excluded_product_types: set[str] | None = None,
        product_type_filter: str | None = None,
    ) -> ResolvedOrderExport:
        allowed_statuses = {
            OrderStatus.NEEDS_REVIEW.value,
            OrderStatus.READY_TO_EXPORT.value,
            OrderStatus.EXPORTED.value,
        }
        if order.status not in allowed_statuses:
            raise ValueError(
                f"Order is not ready to export: status must be {OrderStatus.READY_TO_EXPORT.value} "
                f"or {OrderStatus.NEEDS_REVIEW.value}, "
                f"got {order.status}."
            )
        document_date = date.today()

        if order.client is None or not order.client.client_code:
            raise ValueError("Order is not ready to export: client is not resolved.")
        client_code = order.client.client_code
        client_name = order.client.name

        lines: list[ExportLine] = []
        problems: list[ExportProblemLine] = []
        excluded_product_types = excluded_product_types or set()
        product_type_filter = product_type_filter.strip() if product_type_filter else None
        product_type_filter_key = product_type_filter.casefold() if product_type_filter else None
        for item in sorted(order.items, key=lambda row: row.row_number or 0):
            item_product_type = item.product.product_type if item.product else None
            item_product_type_key = (item_product_type or "").casefold()
            if product_type_filter_key and item_product_type_key != product_type_filter_key:
                continue
            if item.status == OrderItemStatus.SKIPPED.value:
                problems.append(self._problem_line(item, "Строка пропущена"))
                continue
            if item.status == OrderItemStatus.INVALID_QUANTITY.value:
                problems.append(self._problem_line(item, "Ошибка количества"))
                continue
            if item.status != OrderItemStatus.RESOLVED.value or item.product is None:
                problems.append(self._problem_line(item, "Товар не сопоставлен"))
                continue
            if item.product.exclude_from_export:
                problems.append(self._problem_line(item, "Товар исключен из Excel"))
                continue
            if not product_type_filter_key and item_product_type_key and item_product_type_key in excluded_product_types:
                problems.append(self._problem_line(item, "Тип исключен из общего Excel"))
                continue
            item_code = item.product.item_code
            if not item_code:
                problems.append(self._problem_line(item, "Нет номера товара для выгрузки"))
                continue
            if is_short_numeric_item_code(item_code, min_digits=3, max_digits=4):
                problems.append(self._problem_line(item, "Короткий номер товара 3-4 знака"))
                continue
            lines.append(
                ExportLine(
                    item_code=item_code,
                    item_name=self._item_name_for_export(item, item_code),
                    quantity=float(self._export_quantity(item.quantity)),
                    product_type=item.product.product_type,
                    row_number=item.row_number,
                    raw_item_code=item.raw_item_code,
                    raw_barcode=item.raw_barcode,
                    normalized_barcode=item.normalized_barcode,
                    warehouse_no=(
                        item.product.product_type_ref.warehouse_no
                        if item.product.product_type_ref is not None
                        else None
                    ),
                )
            )

        if not lines:
            raise ValueError("Order has no exportable rows. Resolve or keep at least one product before export.")

        sequence = max(sequence_number if sequence_number is not None else order.id * 20, 0)
        return ResolvedOrderExport(
            converter_type=order.converter_type or "unknown",
            client_code=client_code,
            document_date=document_date,
            fiche_no=self._fiche_no_for_sequence(sequence),
            lines=lines,
            client_name=client_name,
            order_id=order.id,
            warehouse_no=self._warehouse_no_from_snapshot(order.parsed_snapshot),
            problems=problems,
            product_type_filter=product_type_filter,
            sequence_number=sequence,
        )

    @staticmethod
    def _next_export_sequence_number(db: Session, branch_id: int | None = None) -> int:
        setting = db.get(AppSetting, ExportService._export_sequence_setting_key(branch_id))
        setting_sequence = ExportService._sequence_from_setting_value(
            setting.value if setting is not None else None
        )
        if setting_sequence is not None:
            return setting_sequence

        sequence = ExportService._next_export_sequence_number_from_events(db, branch_id=branch_id)
        ExportService._store_export_sequence_number(db, sequence, branch_id=branch_id)
        return sequence

    @staticmethod
    def set_next_export_sequence_number(
        db: Session,
        sequence_number: int,
        branch_id: int | None = None,
    ) -> tuple[int, bool]:
        if sequence_number < 0:
            raise ValueError("Export sequence number must be non-negative.")
        if sequence_number > ExportService.MAX_FICHE_SEQUENCE:
            raise ValueError("Export sequence number is too large.")
        current_sequence = ExportService._next_export_sequence_number(db, branch_id=branch_id)
        if sequence_number <= current_sequence:
            return current_sequence, False
        ExportService._store_export_sequence_number(db, sequence_number, branch_id=branch_id)
        return sequence_number, True

    @staticmethod
    def _store_export_sequence_number(
        db: Session,
        sequence_number: int,
        branch_id: int | None = None,
    ) -> None:
        setting = db.get(AppSetting, ExportService._export_sequence_setting_key(branch_id))
        value = {"sequence_number": max(sequence_number, 0)}
        if setting is None:
            setting = AppSetting(key=ExportService._export_sequence_setting_key(branch_id), value=value)
            db.add(setting)
            return
        setting.value = value

    @staticmethod
    def _next_export_sequence_number_from_events(db: Session, branch_id: int | None = None) -> int:
        stmt = select(ProcessingEvent).where(
            ProcessingEvent.event_type == ProcessingEventType.EXPORTED.value,
        )
        if branch_id is not None:
            stmt = stmt.join(Order, ProcessingEvent.order_id == Order.id).where(Order.branch_id == branch_id)
        events = list(
            db.scalars(
                stmt
            )
        )
        if not events:
            return 0

        candidates: list[int] = []
        unknown_legacy_events = 0
        for event in events:
            payload = event.payload or {}
            next_sequence = ExportService._int_payload_value(payload.get("export_sequence_next"))
            if next_sequence is not None:
                candidates.append(next_sequence)
                continue
            start_sequence = ExportService._int_payload_value(payload.get("export_sequence_start"))
            sequence_count = ExportService._int_payload_value(payload.get("export_sequence_count")) or 1
            sequence_step = ExportService._int_payload_value(payload.get("export_sequence_step")) or ExportService.FICHE_SEQUENCE_STEP
            if start_sequence is not None:
                candidates.append(start_sequence + sequence_count * sequence_step)
                continue
            filename_sequence = ExportService._sequence_from_filename(payload.get("filename"))
            if filename_sequence is not None:
                candidates.append(filename_sequence + ExportService.LEGACY_EXPORT_SEQUENCE_STRIDE)
            else:
                unknown_legacy_events += 1
        if unknown_legacy_events:
            candidates.append(unknown_legacy_events * ExportService.LEGACY_EXPORT_SEQUENCE_STRIDE)
        return max(candidates)

    @staticmethod
    def _fiche_no_for_sequence(sequence: int) -> str:
        return f"KA{max(sequence, 0):010d}"

    @staticmethod
    def sequence_from_fiche_no(value: object) -> int | None:
        if not isinstance(value, str):
            return None
        text = value.strip().upper()
        if text.startswith("KA"):
            text = text[2:]
        if not text.isdigit():
            return None
        sequence = int(text)
        return sequence if 0 <= sequence <= ExportService.MAX_FICHE_SEQUENCE else None

    @staticmethod
    def _sequence_from_setting_value(value: object) -> int | None:
        if isinstance(value, dict):
            return ExportService._int_payload_value(value.get("sequence_number"))
        return ExportService._int_payload_value(value)

    @staticmethod
    def _int_payload_value(value: object) -> int | None:
        try:
            return int(value) if value not in (None, "") else None
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _sequence_from_filename(value: object) -> int | None:
        if not isinstance(value, str):
            return None
        match = re.search(r" (\d{4,})(?: [^.]*)?\.xlsx$", value)
        return int(match.group(1)) if match else None

    @staticmethod
    def _export_sequence_count(order: ResolvedOrderExport) -> int:
        product_types = {
            ExportService._display_product_type(line.product_type).casefold()
            for line in order.lines
        }
        return max(len(product_types), 1)

    @staticmethod
    def _warehouse_no_from_snapshot(snapshot: dict | None) -> str | None:
        value = (snapshot or {}).get("warehouse_no")
        return str(value).strip() if value not in (None, "") else None

    @staticmethod
    def _warehouse_no_for_export(value: str | int | None) -> str | int:
        text = str(value).strip() if value is not None else ""
        if not text:
            return 0
        return int(text) if text.isdigit() else text

    @staticmethod
    def _export_quantity(value: Decimal) -> Decimal:
        return value.to_integral_value(rounding=ROUND_FLOOR)

    @staticmethod
    def _warehouse_no_for_product_type_block(
        lines: list[ExportLine],
        fallback: str | int | None,
    ) -> str | int | None:
        for line in lines:
            text = str(line.warehouse_no).strip() if line.warehouse_no is not None else ""
            if text:
                return text
        return fallback

    @staticmethod
    def _excluded_product_types(db: Session, branch_id: int | None = None) -> set[str]:
        stmt = select(ProductTypeExportRule).where(ProductTypeExportRule.exclude_from_export.is_(True))
        if branch_id is not None:
            stmt = stmt.where(ProductTypeExportRule.branch_id == branch_id)
        return {
            rule.product_type.casefold()
            for rule in db.scalars(
                stmt
            )
            if rule.product_type
        }

    @staticmethod
    def _export_sequence_setting_key(branch_id: int | None = None) -> str:
        if branch_id in (None, DEFAULT_BRANCH_ID):
            return ExportService.EXPORT_SEQUENCE_SETTING_KEY
        return f"{ExportService.EXPORT_SEQUENCE_SETTING_KEY}:branch:{branch_id}"

    @staticmethod
    def _sorted_lines(lines: list[ExportLine]) -> list[ExportLine]:
        return sorted(lines, key=lambda line: ExportService._item_code_sort_key(line.item_code))

    @staticmethod
    def _item_code_sort_key(item_code: str) -> tuple[int, int | str, str]:
        code = item_code.strip()
        if code.isdigit():
            return (0, int(code), code)
        return (1, code.casefold(), code)

    @staticmethod
    def _write_export_sheet(worksheet, layout, order: ResolvedOrderExport, lines: list[ExportLine]) -> None:
        block_height = ExportService._export_block_height(layout, lines)
        ExportService._write_export_block(worksheet, layout, order, lines, 0, block_height)

    @staticmethod
    def _write_export_blocks_sheet(
        worksheet,
        layout,
        order: ResolvedOrderExport,
        grouped_lines: list[tuple[str, list[ExportLine]]],
    ) -> None:
        row_offset = 0
        multiple_blocks = len(grouped_lines) > 1
        for index, (_product_type, lines) in enumerate(grouped_lines):
            block_height = ExportService._export_block_height(layout, lines)
            if index > 0:
                ExportService._copy_template_block(worksheet, block_height, row_offset + 1)
            sheet_order = replace(
                order,
                fiche_no=(
                    ExportService._fiche_no_for_sequence(order.sequence_number + index * ExportService.FICHE_SEQUENCE_STEP)
                    if multiple_blocks
                    else order.fiche_no
                ),
                warehouse_no=ExportService._warehouse_no_for_product_type_block(lines, order.warehouse_no),
            )
            ExportService._write_export_block(worksheet, layout, sheet_order, lines, row_offset, block_height)
            row_offset += block_height + 1
        final_used_row = max(row_offset - 1, 0)
        if final_used_row and worksheet.max_row > final_used_row:
            worksheet.delete_rows(final_used_row + 1, worksheet.max_row - final_used_row)

    @staticmethod
    def _export_block_height(layout, lines: list[ExportLine]) -> int:
        return layout.data_start_row + max(len(lines), 1) - 1

    @staticmethod
    def _write_export_block(
        worksheet,
        layout,
        order: ResolvedOrderExport,
        lines: list[ExportLine],
        row_offset: int,
        block_height: int,
    ) -> None:
        client_code_cell = ExportService._offset_coordinate(layout.service_cells["client_code"], row_offset)
        document_date_cell = ExportService._offset_coordinate(layout.service_cells["document_date"], row_offset)
        warehouse_no_cell = ExportService._offset_coordinate(layout.service_cells["warehouse_no"], row_offset)
        fiche_no_cell = ExportService._offset_coordinate(layout.service_cells["fiche_no"], row_offset)

        worksheet[client_code_cell] = order.client_code
        worksheet[document_date_cell] = order.document_date
        if worksheet[document_date_cell].number_format == "General":
            worksheet[document_date_cell].number_format = "dd/mm/yyyy"
        worksheet[warehouse_no_cell] = ExportService._warehouse_no_for_export(order.warehouse_no)
        worksheet[fiche_no_cell] = order.fiche_no

        start_row = layout.data_start_row + row_offset
        row_end = row_offset + block_height
        template_styles = {
            col: copy(worksheet.cell(start_row, col)._style)
            for col in layout.data_style_columns
        }
        template_row_height = worksheet.row_dimensions[start_row].height
        clearable_columns = set(layout.clearable_columns)
        if layout.unit_price_col is not None:
            clearable_columns.add(layout.unit_price_col)

        for row in range(start_row, row_end + 1):
            for col in clearable_columns:
                worksheet.cell(row, col).value = None

        for offset, line in enumerate(ExportService._sorted_lines(lines)):
            row = start_row + offset
            if template_row_height is not None:
                worksheet.row_dimensions[row].height = template_row_height
            for col, style in template_styles.items():
                worksheet.cell(row, col)._style = copy(style)
            worksheet.cell(row, layout.item_code_col).value = line.item_code
            if layout.item_name_col is not None:
                worksheet.cell(row, layout.item_name_col).value = line.item_name or line.item_code
            if layout.unit_col is not None:
                worksheet.cell(row, layout.unit_col).value = line.unit
            worksheet.cell(row, layout.quantity_col).value = line.quantity
            if layout.unit_price_col is not None and line.unit_price is not None:
                worksheet.cell(row, layout.unit_price_col).value = line.unit_price

    @staticmethod
    def _copy_template_block(worksheet, row_count: int, target_start_row: int) -> None:
        for row in range(1, row_count + 1):
            target_row = target_start_row + row - 1
            source_dim = worksheet.row_dimensions[row]
            target_dim = worksheet.row_dimensions[target_row]
            target_dim.height = source_dim.height
            target_dim.hidden = source_dim.hidden
            for col in range(1, worksheet.max_column + 1):
                source = worksheet.cell(row, col)
                target = worksheet.cell(target_row, col)
                target.value = source.value
                if source.has_style:
                    target._style = copy(source._style)
                target.number_format = source.number_format
                target.font = copy(source.font)
                target.fill = copy(source.fill)
                target.border = copy(source.border)
                target.alignment = copy(source.alignment)
                target.protection = copy(source.protection)

        for merged_range in list(worksheet.merged_cells.ranges):
            min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
            if min_row < 1 or max_row > row_count:
                continue
            row_delta = target_start_row - 1
            worksheet.merge_cells(
                start_row=min_row + row_delta,
                start_column=min_col,
                end_row=max_row + row_delta,
                end_column=max_col,
            )

    @staticmethod
    def _offset_coordinate(coordinate: str, row_offset: int) -> str:
        row, column = coordinate_to_tuple(coordinate)
        return f"{openpyxl.utils.get_column_letter(column)}{row + row_offset}"

    @staticmethod
    def _write_problem_sheet(workbook, problems: list[ExportProblemLine]) -> None:
        worksheet = workbook.create_sheet("Проблемные")
        headers = [
            "Строка",
            "Код сети",
            "Штрихкод",
            "Наименование",
            "Тип",
            "Количество",
            "Причина",
        ]
        worksheet.append(headers)
        for problem in problems:
            worksheet.append(
                [
                    problem.row_number,
                    problem.raw_item_code,
                    problem.raw_barcode,
                    problem.raw_name,
                    problem.product_type,
                    problem.quantity,
                    problem.reason,
                ]
            )
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        widths = [10, 22, 18, 48, 18, 14, 32]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        worksheet.freeze_panes = "A2"

    @staticmethod
    def _write_all_products_sheet(
        workbook,
        order: ResolvedOrderExport,
        lines: list[ExportLine],
        problems: list[ExportProblemLine],
    ) -> None:
        worksheet = workbook.create_sheet("Все товары")
        worksheet.append([None, None, None, order.client_name or order.client_code, None, order.document_date])
        worksheet["D1"].font = Font(bold=True)
        worksheet["F1"].font = Font(bold=True)
        worksheet["F1"].number_format = "dd/mm/yyyy"
        headers = [
            "Статус",
            "Строка",
            "Код товара",
            "Наименование",
            "Код сети",
            "Штрихкод",
            "Норм. штрихкод",
            "Тип",
            "Количество",
            "Причина",
        ]
        worksheet.append(headers)
        for line in ExportService._sorted_lines(lines):
            worksheet.append(
                [
                    "Выгружается",
                    line.row_number,
                    line.item_code,
                    line.item_name,
                    line.raw_item_code,
                    line.raw_barcode,
                    line.normalized_barcode,
                    line.product_type,
                    line.quantity,
                    None,
                ]
            )
        for problem in problems:
            worksheet.append(
                [
                    "Проблема",
                    problem.row_number,
                    problem.item_code,
                    problem.item_name or problem.raw_name,
                    problem.raw_item_code,
                    problem.raw_barcode,
                    problem.normalized_barcode,
                    problem.product_type,
                    problem.quantity,
                    problem.reason,
                ]
            )
        for cell in worksheet[2]:
            cell.font = Font(bold=True)
        widths = [14, 10, 24, 52, 22, 18, 18, 18, 14, 34]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        worksheet.freeze_panes = "A3"

    @staticmethod
    def _normalize_workbook_views(workbook) -> None:
        for worksheet in workbook.worksheets:
            for selection in worksheet.sheet_view.selection:
                selection.activeCell = "A1"
                selection.sqref = "A1"

    @staticmethod
    def _problem_line(item: OrderItem, reason: str) -> ExportProblemLine:
        product_item_code = item.product.item_code if item.product is not None else None
        product_name = item.product.name if item.product is not None else None
        return ExportProblemLine(
            row_number=item.row_number,
            item_code=product_item_code,
            item_name=product_name,
            raw_item_code=item.raw_item_code,
            raw_barcode=item.raw_barcode,
            normalized_barcode=item.normalized_barcode,
            raw_name=item.raw_name,
            product_type=item.product.product_type if item.product else None,
            quantity=float(item.quantity) if item.quantity is not None else None,
            reason=reason,
        )

    @staticmethod
    def _display_product_type(value: str | None) -> str:
        text = (value or "").strip()
        return text or "Без типа"

    @staticmethod
    def _unique_sheet_title(existing_titles: list[str], title: str) -> str:
        cleaned = re.sub(r"[\[\]:*?/\\]", "_", title).strip() or "Sheet"
        cleaned = cleaned[:31]
        existing = {item.casefold() for item in existing_titles}
        if cleaned.casefold() not in existing:
            return cleaned
        counter = 2
        while True:
            suffix = f" {counter}"
            candidate = f"{cleaned[:31 - len(suffix)]}{suffix}"
            if candidate.casefold() not in existing:
                return candidate
            counter += 1

    @staticmethod
    def _filename_safe_text(value: str, fallback: str) -> str:
        text = sanitize_filename_part(value, fallback).replace("_", " ")
        return re.sub(r"\s+", " ", text).strip() or fallback

    @staticmethod
    def _item_name_for_export(item: OrderItem, fallback_code: str) -> str:
        product_name = (item.product.name or "").strip() if item.product is not None else ""
        if product_name:
            return product_name
        raw_name = (item.raw_name or "").strip()
        return raw_name or fallback_code
