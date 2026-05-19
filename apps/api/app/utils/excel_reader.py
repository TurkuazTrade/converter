from __future__ import annotations

import logging
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import openpyxl

logger = logging.getLogger(__name__)

SUPPORTED_EXCEL_EXTENSIONS = {".xlsx", ".xls"}


class ExcelReadError(ValueError):
    pass


@dataclass(slots=True)
class SheetData:
    name: str
    rows: list[list[Any]]
    hidden_rows: set[int] = field(default_factory=set)

    def visible_rows(self) -> list[tuple[int, list[Any]]]:
        return [
            (index, row)
            for index, row in enumerate(self.rows, start=1)
            if index not in self.hidden_rows
        ]


def read_workbook(path: str | Path, data_only: bool = True) -> list[SheetData]:
    workbook_path = Path(path)
    suffix = workbook_path.suffix.lower()
    if suffix not in SUPPORTED_EXCEL_EXTENSIONS:
        raise ExcelReadError(f"Unsupported file extension: {suffix}")
    if not workbook_path.exists():
        raise ExcelReadError(f"File not found: {workbook_path}")

    try:
        if suffix == ".xlsx":
            return _read_xlsx(workbook_path, data_only=data_only)
        return _read_xls(workbook_path)
    except ExcelReadError:
        raise
    except PermissionError as exc:
        raise ExcelReadError("Cannot read file. It may be open in Excel.") from exc
    except Exception as exc:
        logger.exception("Failed to read workbook: %s", workbook_path)
        raise ExcelReadError(f"Failed to read Excel file: {exc}") from exc


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace("\u00a0", " ").strip()
    if isinstance(value, datetime | date):
        return value
    return value


def _read_xlsx(path: Path, data_only: bool) -> list[SheetData]:
    workbook = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
    try:
        sheets: list[SheetData] = []
        for worksheet in workbook.worksheets:
            hidden_rows = _read_hidden_rows(path, worksheet.title)
            rows = [
                [_clean_cell(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            sheets.append(SheetData(name=worksheet.title, rows=rows, hidden_rows=hidden_rows))
        return sheets
    finally:
        workbook.close()


def _read_hidden_rows(path: Path, sheet_name: str) -> set[int]:
    try:
        with zipfile.ZipFile(path) as archive:
            sheet_xml = _sheet_xml_path(archive, sheet_name)
            if not sheet_xml:
                return set()
            hidden_rows: set[int] = set()
            with archive.open(sheet_xml) as handle:
                for _, element in ElementTree.iterparse(handle, events=("end",)):
                    if element.tag.endswith("row") and element.attrib.get("hidden") == "1":
                        row_number = element.attrib.get("r")
                        if row_number and row_number.isdigit():
                            hidden_rows.add(int(row_number))
                    element.clear()
            return hidden_rows
    except Exception:
        logger.debug("Could not read hidden rows for %s:%s", path, sheet_name, exc_info=True)
        return set()


def _sheet_xml_path(archive: zipfile.ZipFile, sheet_name: str) -> str | None:
    workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rels = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in rels_root
        if "Id" in rel.attrib and "Target" in rel.attrib
    }
    for sheet in workbook_root.iter():
        if not sheet.tag.endswith("sheet") or sheet.attrib.get("name") != sheet_name:
            continue
        relation_id = sheet.attrib.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        )
        target = rels.get(relation_id or "")
        if not target:
            return None
        target = target.lstrip("/")
        return target if target.startswith("xl/") else f"xl/{target}"
    return None


def _read_xls(path: Path) -> list[SheetData]:
    try:
        import xlrd
    except ImportError as exc:
        raise ExcelReadError("Legacy .xls support requires xlrd.") from exc

    workbook = xlrd.open_workbook(path, formatting_info=True)
    sheets: list[SheetData] = []
    for sheet in workbook.sheets():
        rows: list[list[Any]] = []
        hidden_rows: set[int] = set()
        for zero_index in range(sheet.nrows):
            info = sheet.rowinfo_map.get(zero_index)
            if info and getattr(info, "hidden", 0):
                hidden_rows.add(zero_index + 1)
            values: list[Any] = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(zero_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                    except Exception:
                        logger.debug(
                            "Could not parse xls date at row=%s col=%s",
                            zero_index,
                            col_index,
                        )
                values.append(_clean_cell(value))
            rows.append(values)
        sheets.append(SheetData(name=sheet.name, rows=rows, hidden_rows=hidden_rows))
    return sheets
