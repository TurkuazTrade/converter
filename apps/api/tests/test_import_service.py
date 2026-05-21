from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import OrderItemStatus, OrderStatus
from app.models.client import Client
from app.models.mapping import ClientMapping, ProductMapping
from app.models.order import Order, OrderItem
from app.models.product import Product, ProductBarcode, ProductTypeCatalog
from app.services.import_service import ImportService
from app.services.matching_service import MatchingService


@pytest.mark.asyncio
async def test_import_products_reads_piton_convert_sheet(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [["ignored"]],
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "203200105650133015400030"))
    assert product is not None
    assert product.name == ""
    assert result["converter_type"] == "piton"
    assert result["inserted"] == 1
    assert result["mappings_inserted"] == 1
    assert db_session.scalar(select(ProductBarcode).where(ProductBarcode.barcode == "5029053540108")) is not None
    mapping = db_session.scalar(
        select(ProductMapping).where(
            ProductMapping.converter_type == "piton",
            ProductMapping.normalized_barcode == "5029053540108",
            ProductMapping.product_id == product.id,
        )
    )
    assert mapping is not None
    assert mapping.conversion_multiplier == Decimal("1.000")


@pytest.mark.asyncio
async def test_import_products_skips_duplicate_mappings_in_same_file(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "203200105650133015400030"))
    mappings = list(
        db_session.scalars(
            select(ProductMapping).where(
                ProductMapping.converter_type == "piton",
                ProductMapping.normalized_barcode == "5029053540108",
                ProductMapping.product_id == product.id,
            )
        )
    )
    assert result["inserted"] == 1
    assert result["updated"] == 1
    assert result["mappings_inserted"] == 1
    assert len(mappings) == 1


@pytest.mark.asyncio
async def test_import_products_updates_existing_item_code_without_case_sensitivity(db_session: Session) -> None:
    product = Product(item_code="ERP-Case-1", name="Old name", is_active=True)
    db_session.add(product)
    db_session.flush()
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "Name", "Conv. Quantity"],
                ["erp-case-1", "New name", 2],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    products = list(db_session.scalars(select(Product).where(Product.item_code.ilike("erp-case-1"))))
    assert result["inserted"] == 0
    assert result["updated"] == 1
    assert len(products) == 1
    assert product.name == "New name"
    assert product.conversion_multiplier == Decimal("2.000")


@pytest.mark.asyncio
async def test_import_products_prefers_item_code_and_moves_conflicting_barcode(
    db_session: Session,
) -> None:
    wrong_product = Product(item_code="ERP-WRONG", name="Wrong product", is_active=True)
    correct_product = Product(item_code="ERP-CORRECT", name="Old name", is_active=True)
    db_session.add_all([wrong_product, correct_product])
    db_session.flush()
    wrong_barcode = ProductBarcode(
        product_id=wrong_product.id,
        barcode="1234567890123",
        source="import",
        is_primary=True,
        is_active=True,
    )
    db_session.add(wrong_barcode)
    db_session.flush()
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Name"],
                ["ERP-CORRECT", "1234567890123", "Correct product"],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    active_barcode = db_session.scalar(
        select(ProductBarcode).where(
            ProductBarcode.barcode == "1234567890123",
            ProductBarcode.is_active.is_(True),
        )
    )
    assert result["inserted"] == 0
    assert result["updated"] == 1
    assert correct_product.name == "Correct product"
    assert wrong_barcode.is_active is False
    assert active_barcode is not None
    assert active_barcode.product_id == correct_product.id


@pytest.mark.asyncio
async def test_import_products_does_not_let_weak_numeric_code_steal_barcode_owner(
    db_session: Session,
) -> None:
    code_product = Product(item_code="202", name="Different product", is_active=True)
    barcode_product = Product(item_code=None, name="Old barcode product", is_active=True)
    db_session.add_all([code_product, barcode_product])
    db_session.flush()
    db_session.add(
        ProductBarcode(
            product_id=barcode_product.id,
            barcode="4823077624285",
            source="import",
            is_primary=True,
            is_active=True,
        )
    )
    db_session.flush()
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Name"],
                ["202", "4823077624285", "Roshen Chocolateria"],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    assert result["inserted"] == 0
    assert result["updated"] == 1
    assert code_product.name == "Different product"
    assert barcode_product.name == "Roshen Chocolateria"


@pytest.mark.asyncio
async def test_import_products_keeps_real_human_name(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Name", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", "Мыло Dalan огурец 150г", 1],
            ],
        },
    )

    await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "203200105650133015400030"))
    assert product is not None
    assert product.name == "Мыло Dalan огурец 150г"


@pytest.mark.asyncio
async def test_import_products_reads_catalog_fields_from_product_sheets(db_session: Session) -> None:
    upload = _upload_workbook(
        "база данных по товарам сети Азия Ритейл.xlsx",
        {
            "food": [
                [
                    "Наименование",
                    "Номер товара",
                    "Код обмена",
                    "Штрихкод",
                    "Артикул",
                    "Остаток",
                    "Торговая Марка",
                    "Бренд",
                ],
                [
                    "Печенье Lotte Choco Pie 336г",
                    "201",
                    "EX-1",
                    "4607176441086",
                    "ART-1",
                    "12",
                    "LOTTE",
                    "LOTTE",
                ],
            ],
            "nonfood": [
                [
                    "Наименование",
                    "Номер товара",
                    "Код обмена",
                    "Штрихкод",
                    "Артикул",
                    "Остаток",
                    "Торговая Марка",
                    "Бренд",
                ],
                [
                    "Мыло Dalan Family 5*75гр",
                    "203150105380107012100020",
                    "",
                    "8690529142002",
                    "",
                    "",
                    "DALAN",
                    "DALAN",
                ],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload, converter_type="asia_retail")
    db_session.flush()

    food_product = db_session.scalar(select(Product).where(Product.item_code == "201"))
    nonfood_product = db_session.scalar(
        select(Product).where(Product.item_code == "203150105380107012100020")
    )
    assert result["inserted"] == 2
    assert food_product is not None
    assert food_product.exchange_code == "EX-1"
    assert food_product.article == "ART-1"
    assert food_product.stock == "12"
    assert food_product.trade_mark == "LOTTE"
    assert food_product.brand == "LOTTE"
    assert food_product.product_type == "food"
    assert nonfood_product is not None
    assert nonfood_product.trade_mark == "DALAN"
    assert nonfood_product.brand == "DALAN"
    assert nonfood_product.product_type == "nonfood"


@pytest.mark.asyncio
async def test_import_products_does_not_use_generic_sheet_name_as_product_type(
    db_session: Session,
) -> None:
    upload = _upload_workbook(
        "PRODUCTS.xlsx",
        {
            "Лист1": [
                ["Наименование", "Номер товара", "Штрихкод"],
                ["Печенье Roshen", "201082060195409891300054", "4823077636332"],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "201082060195409891300054"))
    assert result["inserted"] == 1
    assert product is not None
    assert product.product_type is None
    assert (
        db_session.scalar(
            select(ProductTypeCatalog).where(ProductTypeCatalog.normalized_name == "лист1")
        )
        is None
    )


@pytest.mark.asyncio
async def test_import_products_reads_umai_nonfood_fallback_columns(db_session: Session) -> None:
    upload = _upload_workbook(
        "база данных по товарам сети Умай.xlsx",
        {
            "nonfood": [
                ["наименование товара", "длинный код", None, "штрих код", None, None, None],
                [
                    "Крем д/рук Lure витаминный 75мл",
                    "202",
                    None,
                    "4751023296012",
                    "COTTON CLUB",
                    "COTTON CLUB",
                    "#N/A",
                ],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload, converter_type="asia_retail")
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "202"))
    assert result["inserted"] == 1
    assert product is not None
    assert product.trade_mark == "COTTON CLUB"
    assert product.brand == "COTTON CLUB"
    assert product.product_type == "nonfood"


@pytest.mark.asyncio
async def test_import_products_reads_globys_convert_item_code_mapping(db_session: Session) -> None:
    upload = _upload_workbook(
        "GLOBYS CONVERT new.xlsx",
        {
            "convert": [
                ["а", "Client Stock Code", "Conv. Quantity", "Turkuaz Alter 1. Stock Code"],
                ["201082060295409691780001", "Ц0154218", 12, None],
            ]
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "201082060295409691780001"))
    assert product is not None
    assert product.name == ""
    assert result["converter_type"] == "globus"
    assert result["mappings_inserted"] == 1
    mapping = db_session.scalar(
        select(ProductMapping).where(
            ProductMapping.converter_type == "globus",
            ProductMapping.normalized_item_code == "ц0154218",
            ProductMapping.product_id == product.id,
        )
    )
    assert mapping is not None
    assert mapping.conversion_multiplier == Decimal("12.000")

    order = Order(converter_type="globus", status=OrderStatus.PROCESSING.value, parsed_snapshot={})
    item = OrderItem(
        raw_barcode=None,
        normalized_barcode=None,
        raw_item_code="Ц0154218",
        item_code="Ц0154218",
        raw_name="Network product",
        quantity=Decimal("2"),
        row_number=8,
        status=OrderItemStatus.UNRESOLVED.value,
    )
    order.items.append(item)
    db_session.add(order)
    db_session.flush()

    MatchingService(db_session).match_order(order.id)

    assert item.product_id == product.id
    assert item.item_code == "201082060295409691780001"
    assert item.source_quantity == Decimal("2.000")
    assert item.conversion_multiplier == Decimal("12.000")
    assert item.quantity == Decimal("24.000")
    assert item.status == OrderItemStatus.RESOLVED.value


@pytest.mark.asyncio
async def test_import_products_allows_missing_item_code(db_session: Session) -> None:
    upload = _upload_workbook(
        "PRODUCTS.xlsx",
        {
            "convert": [
                ["BARCODE", "Name", "Brand"],
                ["5029053540108", "Мыло без номера", "DALAN"],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.name == "Мыло без номера"))
    assert result["inserted"] == 1
    assert result["skipped"] == 0
    assert result["skipped_file"] is None
    assert product is not None
    assert product.item_code is None
    assert product.brand == "DALAN"
    assert db_session.scalar(select(ProductBarcode).where(ProductBarcode.barcode == "5029053540108")) is not None


@pytest.mark.asyncio
async def test_import_products_returns_skipped_rows_workbook(db_session: Session) -> None:
    upload = _upload_workbook(
        "PRODUCTS.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Name", "Brand"],
                [None, None, None, "DALAN"],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)

    assert result["inserted"] == 0
    assert result["skipped"] == 1
    assert result["skipped_file"] is not None
    assert result["skipped_file"]["filename"] == "PRODUCTS_skipped.xlsx"
    assert result["skipped_file"]["content_base64"]


@pytest.mark.asyncio
async def test_import_clients_reads_piton_client_sheet_and_mapping(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [
                ["код клиента панорама", "название клиента панорама", "название клиента питон"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
            ]
        },
    )

    result = await ImportService(db_session).import_clients(upload)
    db_session.flush()

    client = db_session.scalar(select(Client).where(Client.client_code == "120-04-01-04-8811"))
    assert client is not None
    assert client.name == "Азия Ритейл-1"
    assert client.name_2 == "Гипермаркет 01"
    assert result["converter_type"] == "piton"
    assert result["inserted"] == 1
    assert result["mappings_inserted"] == 1
    assert db_session.scalar(
        select(ClientMapping).where(
            ClientMapping.converter_type == "piton",
            ClientMapping.normalized_client_name == "гипермаркет01",
            ClientMapping.client_id == client.id,
        )
    ) is not None


@pytest.mark.asyncio
async def test_import_clients_skips_duplicate_client_codes_in_same_file(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [
                ["код клиента панорама", "название клиента панорама", "название клиента питон"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
            ]
        },
    )

    result = await ImportService(db_session).import_clients(upload)
    db_session.flush()

    clients = list(db_session.scalars(select(Client).where(Client.client_code == "120-04-01-04-8811")))
    mappings = list(
        db_session.scalars(
            select(ClientMapping).where(
                ClientMapping.converter_type == "piton",
                ClientMapping.normalized_client_name == "гипермаркет01",
            )
        )
    )
    assert result["inserted"] == 1
    assert result["updated"] == 1
    assert result["mappings_inserted"] == 1
    assert len(clients) == 1
    assert len(mappings) == 1


@pytest.mark.asyncio
async def test_import_clients_reads_explicit_second_name(db_session: Session) -> None:
    upload = _upload_workbook(
        "CLIENTS.xlsx",
        {
            "clients": [
                ["Client Code", "Client Name", "Client Name 2"],
                ["C-001", "Название 1", "Название 2"],
            ]
        },
    )

    result = await ImportService(db_session).import_clients(upload, converter_type="piton")
    db_session.flush()

    client = db_session.scalar(select(Client).where(Client.client_code == "C-001"))
    assert client is not None
    assert result["inserted"] == 1
    assert client.name == "Название 1"
    assert client.name_2 == "Название 2"


@pytest.mark.asyncio
async def test_import_reference_workbook_imports_convert_and_client_once(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [
                ["код клиента панорама", "название клиента панорама", "название клиента питон"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
            ],
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_reference_workbook(upload)
    db_session.flush()

    assert result["converter_type"] == "piton"
    assert result["products"]["inserted"] == 1
    assert result["products"]["mappings_inserted"] == 1
    assert result["clients"]["inserted"] == 1
    assert result["clients"]["mappings_inserted"] == 1


def _upload_workbook(filename: str, sheets: dict[str, list[list[object]]]) -> UploadFile:
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)
    for title, rows in sheets.items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return UploadFile(filename=filename, file=buffer)
