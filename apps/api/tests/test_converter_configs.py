from __future__ import annotations

import openpyxl

from app.core.enums import FileRole
from app.services.converter_registry_service import ConverterRegistryService
from app.services.order_processing_service import OrderProcessingService
from app.services.storage_service import StoredObject


def test_converter_configs_load() -> None:
    service = ConverterRegistryService()
    configs = service.list_configs()
    types = {config["type"] for config in configs}
    assert types == {
        "piton",
        "narodnyi",
        "globus",
        "spar",
        "dostor",
        "asia_retail",
        "darkstore",
        "alma",
    }
    for config in configs:
        assert config["version"]
        assert service.config_hash(config["type"])


def test_asia_retail_one_c_export_format_parses(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Склад",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "12",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "МЫЛО ТУАЛЕТНОЕ DALAN",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "120-04-1-03-8812",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = ConverterRegistryService().get_converter("asia_retail").parse(path)

    assert parsed.document_no == "ЦБ00309246"
    assert parsed.warehouse_no == "12"
    assert parsed.snapshot()["warehouse_no"] == "12"
    assert parsed.client_hint.raw_name == "Гипермаркет 12"
    assert parsed.client_hint.client_code == "120-04-1-03-8812"
    assert len(parsed.items) == 1
    assert parsed.items[0].raw_item_code == "203150105380107012200024"
    assert parsed.items[0].normalized_barcode == "8690529522897"
    assert parsed.items[0].quantity == 12


def test_asia_retail_skips_short_numeric_item_codes(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "VALID ITEM",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "SHORT ITEM CODE",
            "4607176441079",
            3,
            116,
            348,
            "201",
            "",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = ConverterRegistryService().get_converter("asia_retail").parse(path)

    assert len(parsed.items) == 1
    assert parsed.items[0].raw_name == "VALID ITEM"


def test_darkstore_email_format_uses_filename_as_document_no(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Лист1"
    worksheet["C1"] = "Заказ на Алма Весна"
    worksheet["C2"] = "Адрес: Джаманбаева 8/2"
    worksheet.append([])
    worksheet.append(["ID", "Штрих-код", "Наименование", "Заказ/шт"])
    worksheet.append([10318913, "4870254131401", "Арахис соленый Big Bob 170гр", 5])
    worksheet.append([10326707, "5029053541648", "Бумага туалетная", 0])
    path = tmp_path / "Алма Весна почта.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = ConverterRegistryService().get_converter("darkstore").parse(path)

    assert parsed.document_no == "Алма Весна почта"
    assert parsed.client_hint.raw_name == "Алма Весна"
    assert parsed.client_hint.raw_address == "Джаманбаева 8/2"
    assert len(parsed.items) == 1
    assert parsed.items[0].raw_item_code == "10318913"
    assert parsed.items[0].normalized_barcode == "4870254131401"
    assert parsed.items[0].quantity == 5


def test_alma_named_darkstore_email_format_autodetects_darkstore(tmp_path, db_session) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Лист1"
    worksheet["C1"] = "Заказ на Алма ГУМ"
    worksheet["C2"] = "Адрес: Чуй 92"
    worksheet.append([])
    worksheet.append(["ID", "Штрих-код", "Наименование", "Заказ/шт"])
    worksheet.append([10323232, "4605496001584", "Вермишель Роллтон", 10])
    path = tmp_path / "Алма ГУМ.xlsx"
    workbook.save(path)
    workbook.close()

    detected = OrderProcessingService(db_session)._detect_converter(
        StoredObject(
            original_name="Алма ГУМ.xlsx",
            stored_name="Алма ГУМ.xlsx",
            path=str(path),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size=path.stat().st_size,
            sha256="test",
            file_role=FileRole.SOURCE,
        )
    )

    assert detected == "darkstore"


def test_autodetect_uses_workbook_content_for_ambiguous_filename(tmp_path, db_session) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "МЫЛО ТУАЛЕТНОЕ DALAN",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    detected = OrderProcessingService(db_session)._detect_converter(
        StoredObject(
            original_name="fr-12.xlsx",
            stored_name="fr-12.xlsx",
            path=str(path),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size=path.stat().st_size,
            sha256="test",
            file_role=FileRole.SOURCE,
        )
    )

    assert detected == "asia_retail"
