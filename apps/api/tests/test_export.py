from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy.orm import Session

from app.api.v1.orders import _export_download_info, _mark_export_downloaded
from app.core.enums import OrderItemStatus, OrderStatus
from app.models.client import Client
from app.models.order import Order, OrderItem, ProcessingEvent
from app.models.product import Product, ProductTypeCatalog, ProductTypeExportRule
from app.models.user import User
from app.services.export_service import ExportLine, ExportService, ResolvedOrderExport


TEMPLATE_PATH = Path("../../data/templates/template_zakaz.xlsx").resolve()


def test_export_matches_template_contract() -> None:
    service = ExportService(template_path=TEMPLATE_PATH)
    content = service.build_export_bytes(
        ResolvedOrderExport(
            converter_type="piton",
            client_code="100245",
            document_date=date(2026, 5, 14),
            fiche_no="0000000001",
            lines=[
                ExportLine(item_code="201300090081421071130040", item_name="Test product 1", quantity=4),
                ExportLine(item_code="ABC-2", item_name="Test product 2", quantity=12.5),
            ],
        )
    )

    template = openpyxl.load_workbook(TEMPLATE_PATH, data_only=False)
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    template_sheet = template.active
    sheet = workbook.active

    assert workbook.sheetnames == template.sheetnames
    assert sheet.title == template_sheet.title
    assert [str(item) for item in sheet.merged_cells.ranges] == [
        str(item) for item in template_sheet.merged_cells.ranges
    ]
    assert [key for key, dim in sheet.column_dimensions.items() if dim.hidden] == ["F"]
    assert sheet.column_dimensions["A"].width == template_sheet.column_dimensions["A"].width
    assert sheet.column_dimensions["B"].width == template_sheet.column_dimensions["B"].width
    assert sheet["A1"].value == "CLIENT CODE"
    assert sheet["B1"].value == "100245"
    assert sheet["A2"].value == "DATE"
    assert _cell_date(sheet["B2"].value) == date(2026, 5, 14)
    assert sheet["B2"].number_format == template_sheet["B2"].number_format
    assert sheet["A3"].value == "WH NO"
    assert sheet["B3"].value == 0
    assert sheet["A4"].value == "FICHE NO"
    assert sheet["B4"].value == "0000000001"
    assert sheet["C5"].value == "Unit"
    assert sheet["E5"].value == "Unit Price (Tenge)"
    assert sheet["A6"].value == "201300090081421071130040"
    assert sheet["B6"].value == "Test product 1"
    assert sheet["C6"].value == 1
    assert sheet["D6"].value == 4
    assert sheet["A7"].value == "ABC-2"
    assert sheet["B7"].value == "Test product 2"
    assert sheet["C7"].value == 1
    assert sheet["D7"].value == 12.5
    assert sheet["E6"].value is None
    assert sheet["A6"].border.left.style == template_sheet["A6"].border.left.style
    assert sheet["D6"].fill.fill_type == template_sheet["D6"].fill.fill_type
    assert sheet.sheet_view.selection[0].activeCell == "A1"
    assert sheet.sheet_view.selection[0].sqref == "A1"
    workbook.close()
    template.close()


def test_export_writes_warehouse_no_when_present() -> None:
    service = ExportService(template_path=TEMPLATE_PATH)
    content = service.build_export_bytes(
        ResolvedOrderExport(
            converter_type="piton",
            client_code="100245",
            document_date=date(2026, 5, 14),
            fiche_no="0000000001",
            lines=[
                ExportLine(item_code="201300090081421071130040", item_name="Test product", quantity=4),
            ],
            warehouse_no="12",
        )
    )

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    assert workbook.active["B3"].value == 12
    workbook.close()


def test_export_order_uses_warehouse_no_from_snapshot(db_session: Session) -> None:
    order = _resolved_order(db_session, warehouse_no="7")
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.active["B3"].value == 7
    workbook.close()


def test_export_keeps_unit_price_headers_once_per_product_type_block() -> None:
    service = ExportService(template_path=TEMPLATE_PATH)
    lines = [
        ExportLine(item_code=f"FOOD-{index:02}", item_name=f"Food {index}", quantity=index, product_type="food")
        for index in range(1, 4)
    ]
    lines.extend(
        ExportLine(
            item_code=f"NONFOOD-{index:02}",
            item_name=f"Nonfood {index}",
            quantity=index,
            product_type="nonfood",
        )
        for index in range(1, 24)
    )
    lines.extend(
        ExportLine(item_code=f"ROSHEN-{index:02}", item_name=f"Roshen {index}", quantity=index, product_type="roshen")
        for index in range(1, 4)
    )
    content = service.build_export_bytes(
        ResolvedOrderExport(
            converter_type="piton",
            client_code="100245",
            document_date=date(2026, 5, 14),
            fiche_no="0000000001",
            lines=lines,
        )
    )

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    unit_price_cells = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == "Unit Price (Tenge)"
    ]
    unit_cells = [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == "Unit"
    ]
    assert unit_price_cells == ["E5", "E14", "E43"]
    assert unit_cells == ["C5", "C14", "C43"]
    workbook.close()


def test_export_lines_are_sorted_by_item_code() -> None:
    service = ExportService(template_path=TEMPLATE_PATH)
    content = service.build_export_bytes(
        ResolvedOrderExport(
            converter_type="piton",
            client_code="100245",
            document_date=date(2026, 5, 14),
            fiche_no="0000000001",
            lines=[
                ExportLine(item_code="203150105380107012200140", item_name="Third", quantity=1),
                ExportLine(item_code="201082060195409891300035", item_name="First", quantity=1),
                ExportLine(item_code="201122065180057605540020", item_name="Second", quantity=1),
            ],
        )
    )

    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    sheet = workbook.active
    assert [sheet[f"A{row}"].value for row in range(6, 9)] == [
        "201082060195409891300035",
        "201122065180057605540020",
        "203150105380107012200140",
    ]
    workbook.close()


def test_export_filename_is_operator_friendly() -> None:
    filename = ExportService().build_filename(
        ResolvedOrderExport(
            converter_type="asia_retail",
            client_code="120-31-4-04-2486",
            document_date=date(2026, 5, 14),
            fiche_no="0000000001",
            lines=[],
            order_id=48,
            sequence_number=20,
        )
    )

    assert filename == "AsiaRetail zakaz 2026-05-14 id-48 0020.xlsx"


def test_export_order_uses_resolved_product_and_client(db_session: Session) -> None:
    order = _resolved_order(db_session, product_item_code="ERP-100", raw_barcode="9999999999999")
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(
        db_session,
        order.id,
    )

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    sheet = workbook.active
    assert sheet["B1"].value == "100245"
    assert _cell_date(sheet["B2"].value) == date.today()
    assert sheet["B4"].value == "KA0000000000"
    assert sheet["A6"].value == "ERP-100"
    assert sheet["A6"].value != "9999999999999"
    assert sheet["B6"].value == "Resolved Product"
    assert sheet["B6"].value != "a"
    assert sheet["D6"].value == 4
    assert order.status == OrderStatus.EXPORTED.value
    assert order.export_file_id is None
    assert result.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0000.xlsx"
    workbook.close()


def test_repeated_export_uses_next_global_fiche_sequence(db_session: Session) -> None:
    order = _resolved_order(db_session)
    db_session.flush()

    first = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)
    second = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    first_workbook = openpyxl.load_workbook(BytesIO(first.content), data_only=True)
    second_workbook = openpyxl.load_workbook(BytesIO(second.content), data_only=True)
    assert first.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0000.xlsx"
    assert second.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0005.xlsx"
    assert first_workbook.active["B4"].value == "KA0000000000"
    assert second_workbook.active["B4"].value == "KA0000000005"
    assert (db_session.get(Order, order.id).events[-1].payload or {})["export_sequence_next"] == 10
    first_workbook.close()
    second_workbook.close()


def test_export_sequence_counts_existing_export_events(db_session: Session) -> None:
    order = _resolved_order(db_session)
    db_session.flush()
    db_session.add(
        ProcessingEvent(
            order_id=order.id,
            event_type="exported",
            message="Existing export.",
            payload={},
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert result.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0050.xlsx"
    assert workbook.active["B4"].value == "KA0000000050"
    workbook.close()


def test_export_sequence_is_global_across_orders(db_session: Session) -> None:
    first_order = _resolved_order(
        db_session,
        product_item_code="ERP-FIRST",
        client_code="100245-GLOBAL-1",
    )
    second_order = _resolved_order(
        db_session,
        product_item_code="ERP-SECOND",
        client_code="100245-GLOBAL-2",
    )
    db_session.flush()

    first = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, first_order.id)
    second = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, second_order.id)

    first_workbook = openpyxl.load_workbook(BytesIO(first.content), data_only=True)
    second_workbook = openpyxl.load_workbook(BytesIO(second.content), data_only=True)
    assert first.filename == f"Piton zakaz {date.today().isoformat()} id-{first_order.id} 0000.xlsx"
    assert second.filename == f"Piton zakaz {date.today().isoformat()} id-{second_order.id} 0005.xlsx"
    assert first_workbook.active["B4"].value == "KA0000000000"
    assert second_workbook.active["B4"].value == "KA0000000005"
    first_workbook.close()
    second_workbook.close()


def test_next_export_sequence_skips_all_blocks_from_previous_full_export(db_session: Session) -> None:
    order = _resolved_order(db_session, product_type="food")
    nonfood_product = Product(
        item_code="ERP-NONFOOD",
        name="Nonfood Product",
        product_type="nonfood",
        is_active=True,
    )
    kcc_product = Product(
        item_code="ERP-KCC",
        name="KCC Product",
        product_type="kcc",
        is_active=True,
    )
    db_session.add_all([nonfood_product, kcc_product])
    db_session.flush()
    order.items.extend(
        [
            OrderItem(
                product_id=nonfood_product.id,
                raw_barcode="4600000000001",
                normalized_barcode="4600000000001",
                raw_name="Nonfood Raw Product",
                raw_item_code="RAW-NONFOOD",
                item_code="RAW-NONFOOD",
                quantity=Decimal("6"),
                row_number=7,
                status=OrderItemStatus.RESOLVED.value,
            ),
            OrderItem(
                product_id=kcc_product.id,
                raw_barcode="4600000000002",
                normalized_barcode="4600000000002",
                raw_name="KCC Raw Product",
                raw_item_code="RAW-KCC",
                item_code="RAW-KCC",
                quantity=Decimal("7"),
                row_number=8,
                status=OrderItemStatus.RESOLVED.value,
            ),
        ]
    )
    db_session.flush()

    first = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)
    second = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    first_workbook = openpyxl.load_workbook(BytesIO(first.content), data_only=True)
    second_workbook = openpyxl.load_workbook(BytesIO(second.content), data_only=True)
    assert first.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0000.xlsx"
    assert first_workbook["1"]["B4"].value == "KA0000000000"
    assert first_workbook["1"]["B11"].value == "KA0000000005"
    assert first_workbook["1"]["B18"].value == "KA0000000010"
    assert second.filename == f"Piton zakaz {date.today().isoformat()} id-{order.id} 0015.xlsx"
    assert second_workbook["1"]["B4"].value == "KA0000000015"
    first_workbook.close()
    second_workbook.close()


def test_export_needs_review_order_with_resolved_rows(db_session: Session) -> None:
    order = _resolved_order(db_session, status=OrderStatus.NEEDS_REVIEW.value)
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.active["A6"].value == "ERP-100"
    workbook.close()


def test_export_skips_unresolved_items(db_session: Session) -> None:
    order = _resolved_order(db_session, status=OrderStatus.NEEDS_REVIEW.value)
    order.items.append(
        OrderItem(
            product_id=None,
            raw_barcode="404",
            normalized_barcode="404",
            raw_name="Unknown Product",
            raw_item_code="RAW-404",
            item_code="RAW-404",
            quantity=Decimal("8"),
            row_number=7,
            status=OrderItemStatus.UNRESOLVED.value,
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    sheet = workbook.active
    assert sheet["A6"].value == "ERP-100"
    assert sheet["A7"].value is None
    workbook.close()


def test_export_skips_products_excluded_from_export(db_session: Session) -> None:
    order = _resolved_order(db_session)
    excluded_product = db_session.get(Product, order.items[0].product_id)
    assert excluded_product is not None
    excluded_product.exclude_from_export = True
    included_product = Product(
        item_code="ERP-INCLUDED",
        name="Included Product",
        is_active=True,
    )
    db_session.add(included_product)
    db_session.flush()
    order.items.append(
        OrderItem(
            product_id=included_product.id,
            raw_barcode="4600000000001",
            normalized_barcode="4600000000001",
            raw_name="Included Raw Product",
            raw_item_code="RAW-INCLUDED",
            item_code="RAW-INCLUDED",
            quantity=Decimal("6"),
            row_number=7,
            status=OrderItemStatus.RESOLVED.value,
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    sheet = workbook.active
    assert sheet["A6"].value == "ERP-INCLUDED"
    assert sheet["D6"].value == 6
    assert sheet["A7"].value is None
    workbook.close()


def test_export_skips_products_with_excluded_product_type(db_session: Session) -> None:
    order = _resolved_order(db_session, product_type="food")
    included_product = Product(
        item_code="ERP-INCLUDED",
        name="Included Product",
        product_type="nonfood",
        is_active=True,
    )
    db_session.add(
        ProductTypeExportRule(
            product_type="food",
            exclude_from_export=True,
        )
    )
    db_session.add(included_product)
    db_session.flush()
    order.items.append(
        OrderItem(
            product_id=included_product.id,
            raw_barcode="4600000000001",
            normalized_barcode="4600000000001",
            raw_name="Included Raw Product",
            raw_item_code="RAW-INCLUDED",
            item_code="RAW-INCLUDED",
            quantity=Decimal("6"),
            row_number=7,
            status=OrderItemStatus.RESOLVED.value,
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    sheet = workbook.active
    assert sheet["A6"].value == "ERP-INCLUDED"
    assert sheet["D6"].value == 6
    assert sheet["A7"].value is None
    workbook.close()


def test_export_stacks_product_type_blocks_on_one_sheet_and_adds_problem_sheet(db_session: Session) -> None:
    order = _resolved_order(db_session, product_type="food")
    food_type = ProductTypeCatalog(
        name="food",
        normalized_name="food",
        warehouse_no="1",
        is_active=True,
    )
    nonfood_type = ProductTypeCatalog(
        name="nonfood",
        normalized_name="nonfood",
        warehouse_no="5",
        is_active=True,
    )
    db_session.add_all([food_type, nonfood_type])
    db_session.flush()
    food_product = db_session.get(Product, order.items[0].product_id)
    assert food_product is not None
    food_product.product_type_id = food_type.id
    nonfood_product = Product(
        item_code="ERP-NONFOOD",
        name="Nonfood Product",
        product_type="nonfood",
        product_type_id=nonfood_type.id,
        is_active=True,
    )
    db_session.add(nonfood_product)
    db_session.flush()
    order.items.append(
        OrderItem(
            product_id=nonfood_product.id,
            raw_barcode="4600000000002",
            normalized_barcode="4600000000002",
            raw_name="Nonfood Raw Product",
            raw_item_code="RAW-NONFOOD",
            item_code="RAW-NONFOOD",
            quantity=Decimal("7"),
            row_number=7,
            status=OrderItemStatus.RESOLVED.value,
        )
    )
    order.items.append(
        OrderItem(
            product_id=None,
            raw_barcode="404",
            normalized_barcode="404",
            raw_name="Unknown Product",
            raw_item_code="RAW-404",
            item_code="RAW-404",
            quantity=Decimal("8"),
            row_number=8,
            status=OrderItemStatus.UNRESOLVED.value,
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.sheetnames[-1] == "Проблемные"
    assert workbook.sheetnames == ["1", "Проблемные"]
    sheet = workbook["1"]
    assert sheet["A6"].value == "ERP-100"
    assert sheet["B3"].value == 1
    assert sheet["B4"].value == "KA0000000000"
    assert sheet["A7"].value is None
    assert sheet["A13"].value == "ERP-NONFOOD"
    assert sheet["B10"].value == 5
    assert sheet["B11"].value == "KA0000000005"
    assert workbook["Проблемные"]["B2"].value == "RAW-404"
    assert workbook["Проблемные"]["G2"].value == "Товар не сопоставлен"
    workbook.close()


def test_export_moves_short_numeric_item_codes_to_problem_sheet(db_session: Session) -> None:
    order = _resolved_order(db_session, product_item_code="201", product_type="food")
    included_product = Product(
        item_code="ERP-INCLUDED",
        name="Included Product",
        product_type="food",
        is_active=True,
    )
    db_session.add(included_product)
    db_session.flush()
    order.items.append(
        OrderItem(
            product_id=included_product.id,
            raw_barcode="4600000000003",
            normalized_barcode="4600000000003",
            raw_name="Included Raw Product",
            raw_item_code="RAW-INCLUDED",
            item_code="RAW-INCLUDED",
            quantity=Decimal("6"),
            row_number=7,
            status=OrderItemStatus.RESOLVED.value,
        )
    )
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook["1"]["A6"].value == "ERP-INCLUDED"
    assert workbook["Проблемные"]["B2"].value == "RAW-100"
    assert workbook["Проблемные"]["G2"].value == "Короткий номер товара 3-4 знака"
    workbook.close()


def test_export_by_type_ignores_type_exclusion_rule(db_session: Session) -> None:
    order = _resolved_order(db_session, product_type="food")
    db_session.add(ProductTypeExportRule(product_type="food", exclude_from_export=True))
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(
        db_session,
        order.id,
        product_type="food",
    )

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.sheetnames == ["1"]
    assert workbook["1"]["A6"].value == "ERP-100"
    workbook.close()


def test_export_blocks_unresolved_items(db_session: Session) -> None:
    order = _resolved_order(db_session)
    item = order.items[0]
    item.status = OrderItemStatus.UNRESOLVED.value
    item.product_id = None
    db_session.flush()

    with pytest.raises(ValueError, match="no exportable rows"):
        ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)


def test_exported_order_can_be_regenerated(db_session: Session) -> None:
    order = _resolved_order(db_session, status=OrderStatus.EXPORTED.value)
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(
        db_session,
        order.id,
    )

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.active["B6"].value == "Resolved Product"
    workbook.close()


def test_export_falls_back_to_product_name_when_source_name_missing(db_session: Session) -> None:
    order = _resolved_order(db_session)
    order.items[0].raw_name = ""
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(
        db_session,
        order.id,
    )

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.active["B6"].value == "Resolved Product"
    workbook.close()


def test_export_falls_back_to_source_name_when_product_name_missing(db_session: Session) -> None:
    order = _resolved_order(db_session)
    product = db_session.get(Product, order.items[0].product_id)
    assert product is not None
    product.name = ""
    db_session.flush()

    result = ExportService(template_path=TEMPLATE_PATH).export_order(
        db_session,
        order.id,
    )

    workbook = openpyxl.load_workbook(BytesIO(result.content), data_only=True)
    assert workbook.active["B6"].value == "Raw Product"
    workbook.close()


def test_export_blocks_order_without_client(db_session: Session) -> None:
    order = _resolved_order(db_session)
    order.client_id = None
    db_session.flush()

    with pytest.raises(ValueError, match="client is not resolved"):
        ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)


def test_export_blocks_product_without_item_code(db_session: Session) -> None:
    order = _resolved_order(db_session, product_item_code=None)
    db_session.flush()

    with pytest.raises(ValueError, match="no exportable rows"):
        ExportService(template_path=TEMPLATE_PATH).export_order(db_session, order.id)


def test_export_download_marker_tracks_each_variant(db_session: Session) -> None:
    order = _resolved_order(db_session)
    user = User(email="operator@example.com", hashed_password="hash", full_name="Operator One", role="operator")
    db_session.add(user)
    db_session.flush()

    _mark_export_downloaded(order, "flint", user)
    first_downloaded_at = order.export_downloaded_at
    _mark_export_downloaded(order, "food", user)

    assert first_downloaded_at is not None
    assert set(order.export_downloads or {}) == {"flint", "food"}
    assert order.export_downloads["flint"] == {
        "downloaded_at": first_downloaded_at.isoformat(),
        "user_id": user.id,
        "user_name": "Operator One",
    }
    assert order.export_downloads["food"] == {
        "downloaded_at": order.export_downloaded_at.isoformat(),
        "user_id": user.id,
        "user_name": "Operator One",
    }
    assert _export_download_info(order, "flint")["user_name"] == "Operator One"


def _resolved_order(
    db_session: Session,
    *,
    status: str = OrderStatus.READY_TO_EXPORT.value,
    product_item_code: str | None = "ERP-100",
    raw_barcode: str = "4600000000000",
    client_code: str = "100245",
    product_type: str | None = None,
    warehouse_no: str | None = None,
) -> Order:
    client = Client(
        client_code=client_code,
        name="Resolved Client",
        normalized_name="resolvedclient",
        is_active=True,
    )
    product = Product(
        item_code=product_item_code,
        name="Resolved Product",
        product_type=product_type,
        is_active=True,
    )
    db_session.add_all([client, product])
    db_session.flush()

    order = Order(
        order_number="0000000001",
        converter_type="piton",
        converter_version="1.0.0",
        client_id=client.id,
        status=status,
        parsed_snapshot={"document_date": "2026-05-14", "warehouse_no": warehouse_no},
    )
    order.items.append(
        OrderItem(
            product_id=product.id,
            raw_barcode=raw_barcode,
            normalized_barcode=raw_barcode,
            raw_name="Raw Product",
            raw_item_code="RAW-100",
            item_code="RAW-100",
            quantity=Decimal("4"),
            row_number=6,
            status=OrderItemStatus.RESOLVED.value,
        )
    )
    db_session.add(order)
    return order


def _cell_date(value: object) -> date:
    if hasattr(value, "date"):
        return value.date()
    return value
