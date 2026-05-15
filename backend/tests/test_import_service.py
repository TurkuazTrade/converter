from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.enums import OrderItemStatus, OrderStatus
from app.models.client import Client
from app.models.mapping import ClientMapping, ProductMapping
from app.models.order import Order, OrderItem
from app.models.product import Product, ProductBarcode
from app.services.import_service import ImportService
from app.services.matching_service import MatchingService


@pytest.mark.asyncio
async def test_import_products_reads_piton_convert_sheet(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [["ignored"]],
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "203200105650133015400030"))
    assert product is not None
    assert result["converter_type"] == "piton"
    assert result["inserted"] == 1
    assert result["mappings_inserted"] == 1
    assert db_session.scalar(select(ProductBarcode).where(ProductBarcode.barcode == "5029053540108")) is not None
    assert db_session.scalar(
        select(ProductMapping).where(
            ProductMapping.converter_type == "piton",
            ProductMapping.normalized_barcode == "5029053540108",
            ProductMapping.product_id == product.id,
        )
    ) is not None


@pytest.mark.asyncio
async def test_import_products_skips_duplicate_mappings_in_same_file(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "203200105650133015400030"))
    mappings = list(
        db_session.scalars(
            select(ProductMapping).where(
                ProductMapping.converter_type == "piton",
                ProductMapping.normalized_barcode == "5029053540108",
                ProductMapping.product_id == product.id,
            )
        )
    )
    assert result["inserted"] == 1
    assert result["updated"] == 1
    assert result["mappings_inserted"] == 1
    assert len(mappings) == 1


@pytest.mark.asyncio
async def test_import_products_reads_globys_convert_item_code_mapping(db_session: Session) -> None:
    upload = _upload_workbook(
        "GLOBYS CONVERT new.xlsx",
        {
            "convert": [
                ["а", "Client Stock Code", "Conv. Quantity", "Turkuaz Alter 1. Stock Code"],
                ["201082060295409691780001", "Ц0154218", 1, None],
            ]
        },
    )

    result = await ImportService(db_session).import_products(upload)
    db_session.flush()

    product = db_session.scalar(select(Product).where(Product.item_code == "201082060295409691780001"))
    assert product is not None
    assert result["converter_type"] == "globus"
    assert result["mappings_inserted"] == 1
    assert db_session.scalar(
        select(ProductMapping).where(
            ProductMapping.converter_type == "globus",
            ProductMapping.normalized_item_code == "ц0154218",
            ProductMapping.product_id == product.id,
        )
    ) is not None

    order = Order(converter_type="globus", status=OrderStatus.PROCESSING.value, parsed_snapshot={})
    item = OrderItem(
        raw_barcode=None,
        normalized_barcode=None,
        raw_item_code="Ц0154218",
        item_code="Ц0154218",
        raw_name="Network product",
        quantity=Decimal("2"),
        row_number=8,
        status=OrderItemStatus.UNRESOLVED.value,
    )
    order.items.append(item)
    db_session.add(order)
    db_session.flush()

    MatchingService(db_session).match_order(order.id)

    assert item.product_id == product.id
    assert item.item_code == "201082060295409691780001"
    assert item.status == OrderItemStatus.RESOLVED.value


@pytest.mark.asyncio
async def test_import_clients_reads_piton_client_sheet_and_mapping(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [
                ["код клиента панорама", "название клиента панорама", "название клиента питон"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
            ]
        },
    )

    result = await ImportService(db_session).import_clients(upload)
    db_session.flush()

    client = db_session.scalar(select(Client).where(Client.client_code == "120-04-01-04-8811"))
    assert client is not None
    assert result["converter_type"] == "piton"
    assert result["inserted"] == 1
    assert result["mappings_inserted"] == 1
    assert db_session.scalar(
        select(ClientMapping).where(
            ClientMapping.converter_type == "piton",
            ClientMapping.normalized_client_name == "гипермаркет01",
            ClientMapping.client_id == client.id,
        )
    ) is not None


@pytest.mark.asyncio
async def test_import_reference_workbook_imports_convert_and_client_once(db_session: Session) -> None:
    upload = _upload_workbook(
        "PITON CONVERT.xlsx",
        {
            "client": [
                ["код клиента панорама", "название клиента панорама", "название клиента питон"],
                ["120-04-01-04-8811", "Азия Ритейл-1", "Гипермаркет 01"],
            ],
            "convert": [
                ["SKU_NO", "BARCODE", "Conv. Quantity"],
                ["203200105650133015400030", "5029053540108", 1],
            ],
        },
    )

    result = await ImportService(db_session).import_reference_workbook(upload)
    db_session.flush()

    assert result["converter_type"] == "piton"
    assert result["products"]["inserted"] == 1
    assert result["products"]["mappings_inserted"] == 1
    assert result["clients"]["inserted"] == 1
    assert result["clients"]["mappings_inserted"] == 1


def _upload_workbook(filename: str, sheets: dict[str, list[list[object]]]) -> UploadFile:
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)
    for title, rows in sheets.items():
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return UploadFile(filename=filename, file=buffer)
