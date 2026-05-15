from __future__ import annotations

import openpyxl

from app.core.enums import FileRole
from app.services.converter_registry_service import ConverterRegistryService
from app.services.order_processing_service import OrderProcessingService
from app.services.storage_service import StoredObject


def test_converter_configs_load() -> None:
    service = ConverterRegistryService()
    configs = service.list_configs()
    types = {config["type"] for config in configs}
    assert types == {
        "piton",
        "narodnyi",
        "globus",
        "spar",
        "dostor",
        "asia_retail",
        "darkstore",
        "alma",
    }
    for config in configs:
        assert config["version"]
        assert service.config_hash(config["type"])


def test_asia_retail_one_c_export_format_parses(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "МЫЛО ТУАЛЕТНОЕ DALAN",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = ConverterRegistryService().get_converter("asia_retail").parse(path)

    assert parsed.document_no == "ЦБ00309246"
    assert parsed.client_hint.raw_name == "Гипермаркет 12"
    assert len(parsed.items) == 1
    assert parsed.items[0].raw_item_code == "203150105380107012200024"
    assert parsed.items[0].normalized_barcode == "8690529522897"
    assert parsed.items[0].quantity == 12


def test_asia_retail_skips_short_numeric_item_codes(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "VALID ITEM",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "SHORT ITEM CODE",
            "4607176441079",
            3,
            116,
            348,
            "201",
            "",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = ConverterRegistryService().get_converter("asia_retail").parse(path)

    assert len(parsed.items) == 1
    assert parsed.items[0].raw_name == "VALID ITEM"


def test_autodetect_uses_workbook_content_for_ambiguous_filename(tmp_path, db_session) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Заявка"
    worksheet.append(
        [
            "Вид документа",
            "Дата",
            "Номер",
            "Код обмена контрагент",
            "Контрагент",
            "Подразделение",
            "Код товара",
            "Товар",
            "Штрихкод",
            "Количество",
            "Цена",
            "Сумма",
            "Номер товара",
            "Номер контрагент",
        ]
    )
    worksheet.append(
        [
            "Заявки",
            "15.05.2026",
            "ЦБ00309246",
            "",
            'ЗАО "Азия Ритейл"',
            "Гипермаркет 12",
            "",
            "МЫЛО ТУАЛЕТНОЕ DALAN",
            "8690529522897",
            12,
            58,
            696,
            "203150105380107012200024",
            "",
        ]
    )
    path = tmp_path / "fr-12.xlsx"
    workbook.save(path)
    workbook.close()

    detected = OrderProcessingService(db_session)._detect_converter(
        StoredObject(
            original_name="fr-12.xlsx",
            stored_name="fr-12.xlsx",
            path=str(path),
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            size=path.stat().st_size,
            sha256="test",
            file_role=FileRole.SOURCE,
        )
    )

    assert detected == "asia_retail"
